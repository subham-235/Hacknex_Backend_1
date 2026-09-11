"""Structural document QA when a page renderer is unavailable."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
path = Path(__file__).with_name("Suraksha-Agent-Workflow.docx")
with ZipFile(path) as archive:
    assert archive.testzip() is None
doc = Document(path)
section = doc.sections[0]
assert section.page_width == Inches(8.5) and section.page_height == Inches(11)
assert section.left_margin == Inches(1) and section.right_margin == Inches(1)
assert doc.styles["Normal"].font.size == Pt(11)
assert doc.styles["Normal"].paragraph_format.line_spacing == 1.25
assert doc.styles["Heading 1"].paragraph_format.space_before == Pt(18)
assert sum(p.style.name == "Title" for p in doc.paragraphs) == 1
assert sum(p.style.name == "List Number" for p in doc.paragraphs) == 8
for p in doc.paragraphs:
    if p.style.name == "List Number": assert p._p.pPr.find(qn("w:numPr")) is not None
for table in doc.tables:
    assert table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
    assert table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")) == "120"
    assert sum(int(c.get(qn("w:w"))) for c in table._tbl.tblGrid.gridCol_lst) == 9360
    for row in table.rows:
        assert sum(int(c._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w"))) for c in row.cells) == 9360
text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for required in ["AGENT_MODE", "agent:worker", "webhooks/incoming", "review_required", "unknown", "No live Gemini", "getEmergencyContext"]:
    assert required in text, required
assert "TODO" not in text and "TBD" not in text
print(f"Structural QA passed: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, real numbered workflow, explicit geometry. Visual QA not performed.")
