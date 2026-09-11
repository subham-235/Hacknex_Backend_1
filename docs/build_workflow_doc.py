"""Generate the Word guide from agent-workflow.md using python-docx.

Design: compact_reference_guide + memo_masthead.
Named overrides: Title 26pt; Subtitle 12pt; Code Consolas 9pt / 1.05;
table text 9pt / 1.05 for API and configuration reference tables.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

def style(name, size, before=0, after=6, color="222222", line=1.25):
    s = doc.styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.line_spacing = line
    s.paragraph_format.widow_control = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return s

style("Normal", 11)
style("Title", 26, after=6, color="1F4D78")
style("Subtitle", 12, after=12, color="555555")
for name, size, before, after, color in [("Heading 1",16,18,10,"2E74B5"),("Heading 2",13,14,7,"2E74B5"),("Heading 3",12,10,5,"1F4D78")]:
    style(name, size, before, after, color).paragraph_format.keep_with_next = True
style("Header", 9, after=0, color="666666", line=1)
style("Footer", 9, after=0, color="666666", line=1)
style("List Number", 11, after=4)
style("List Bullet", 11, after=4)
from docx.enum.style import WD_STYLE_TYPE
doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
code_style = style("Code", 9, after=2, line=1.05)
code_style.font.name = "Consolas"
doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
style("Table Text", 9, after=3, line=1.05)

# Explicit real numbering, with aligned wrapped lines.
numbering = doc.part.numbering_part.element
abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), "42")
lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
for tag, val in [("start","1"),("numFmt","decimal"),("lvlText","%1."),("lvlJc","left")]:
    el = OxmlElement("w:" + tag); el.set(qn("w:val"), val); lvl.append(el)
pp = OxmlElement("w:pPr")
ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "271"); pp.append(ind)
lvl.append(pp); abstract.append(lvl); numbering.append(abstract)
num = OxmlElement("w:num"); num.set(qn("w:numId"), "42")
aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), "42"); num.append(aid); numbering.append(num)

header = section.header.paragraphs[0]
header.text = "SURAKSHA  /  ENGINEERING GUIDE"
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Suraksha SOS Agent  |  ")
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

def add_table(rows):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    widths = [4900, 4460]
    for _ in rows[1:]: table.add_row()
    pr = table._tbl.tblPr
    tw = pr.find(qn("w:tblW")); tw.set(qn("w:w"), "9360"); tw.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd"); indent.set(qn("w:w"), "120"); indent.set(qn("w:type"), "dxa"); pr.append(indent)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in [("top",80),("bottom",80),("start",120),("end",120)]:
        el = OxmlElement("w:"+edge); el.set(qn("w:w"),str(value)); el.set(qn("w:type"),"dxa"); margins.append(el)
    pr.append(margins)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top","left","bottom","right","insideH","insideV"]:
        el = OxmlElement("w:"+edge); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"4"); el.set(qn("w:color"),"CBD5E1"); borders.append(el)
    pr.append(borders)
    for grid, width in zip(table._tbl.tblGrid.gridCol_lst, widths): grid.set(qn("w:w"),str(width))
    for index, (row, values) in enumerate(zip(table.rows, rows)):
        cant_split = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cant_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader"); row._tr.get_or_add_trPr().append(repeat)
        for cell, text, width in zip(row.cells, values, widths):
            cell.width = Inches(width/1440)
            cell._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"),str(width))
            cell.text = text
            p = cell.paragraphs[0]; p.style = doc.styles["Table Text"]
            if index == 0:
                p.runs[0].bold = True
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"),"E8EEF5"); cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = (ROOT / "agent-workflow.md").read_text(encoding="utf-8").splitlines()
i = 0; in_code = False; first = True
while i < len(lines):
    line = lines[i].strip(); i += 1
    if line == "```": in_code = not in_code; continue
    if in_code: doc.add_paragraph(line, style="Code"); continue
    if not line: continue
    if line == "---": doc.add_page_break(); continue
    if line.startswith("| "):
        rows = [[v.strip() for v in line.strip("|").split("|")]]
        while i < len(lines) and lines[i].startswith("| "):
            rows.append([v.strip() for v in lines[i].strip("|").split("|")]); i += 1
        add_table(rows); continue
    if line.startswith("# "):
        p = doc.add_paragraph(line[2:], style="Title" if first else "Heading 1")
        if first:
            p.paragraph_format.keep_with_next = True
            first = False
        continue
    if line.startswith("## "): doc.add_paragraph(line[3:], style="Heading 2"); continue
    if re.match(r"^\d+\. ", line):
        p = doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        props = p._p.get_or_add_pPr(); n = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl"); level.set(qn("w:val"),"0"); n.append(level)
        ident = OxmlElement("w:numId"); ident.set(qn("w:val"),"42"); n.append(ident); props.append(n)
        continue
    doc.add_paragraph(line, style="Subtitle" if line == "Workflow and operations guide" else "Normal")

doc.core_properties.title = "Suraksha SOS Agent: Workflow and Operations Guide"
doc.core_properties.subject = "Implemented agent workflow, configuration, API and failure handling"
doc.core_properties.author = "Suraksha"
doc.save(ROOT / "Suraksha-Agent-Workflow.docx")
print(ROOT / "Suraksha-Agent-Workflow.docx")
